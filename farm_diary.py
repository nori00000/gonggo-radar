#!/usr/bin/env python3
"""
agrion.kr 영농일지 일괄등록 엑셀 생성기

사용법:
  python3 farm_diary.py --start 2026-01-02 --end 2026-03-24 --freq 6
  python3 farm_diary.py --start 2026-04-01 --end 2026-06-30 --freq 3
  python3 farm_diary.py --start 2026-01-02 --end 2026-03-24 --template ~/Downloads/영농일지+일괄등록.xlsx
"""

import argparse
import json
import os
import sys
import glob
import copy
import ssl
import random
from datetime import datetime, timedelta
from urllib.request import urlopen, Request
from urllib.parse import urlencode

# ============================================================
# 설정
# ============================================================

LATITUDE = 37.6762   # 고양시 일산서구 대화동
LONGITUDE = 126.7550

# 2026년 한국 공휴일
HOLIDAYS_2026 = {
    '2026-01-01': '신정',
    '2026-02-16': '설날 연휴',
    '2026-02-17': '설날',
    '2026-02-18': '설날 연휴',
    '2026-03-01': '삼일절',
    '2026-03-02': '삼일절 대체공휴일',
    '2026-05-05': '어린이날',
    '2026-05-24': '부처님오신날',
    '2026-06-06': '현충일',
    '2026-08-15': '광복절',
    '2026-09-24': '추석 연휴',
    '2026-09-25': '추석',
    '2026-09-26': '추석 연휴',
    '2026-10-03': '개천절',
    '2026-10-09': '한글날',
    '2026-12-25': '성탄절',
}

# 작업 패턴 (더 다양한 내용)
PATTERNS = [
    # ---- 파종 관련 (A) ----
    {'step': '파종', 'content': '이끼 배지 정리 및 잔여물 제거\n기반 토양 수분 상태 확인 후 정비\n파종 구역 구획 정리'},
    {'step': '파종', 'content': '이끼 포자 파종 작업 실시\n파종 후 분무 관수 실시\n배지 표면 상태 점검 및 정리'},
    {'step': '파종', 'content': '이끼 모종 배치 및 안착 작업\n배지별 파종 밀도 조절\n파종 후 초기 수분 공급'},
    {'step': '파종', 'content': '이끼 파종 구역 추가 확장 작업\n새 배지 투입 후 파종 준비\n기존 파종 구역 활착 상태 확인'},

    # ---- 물주기 관련 (B) ----
    {'step': '물주기', 'content': '이끼 관수작업 실시 (분무식)\n포트별 수분 상태 개별 확인\n습도 점검 후 조절 작업'},
    {'step': '물주기', 'content': '이끼 관수 후 배수 상태 확인\n과습 구역 배수로 정비\n습도 유지 관리 및 기록'},
    {'step': '물주기', 'content': '이끼 재배상 전체 관수 실시\n분무기 노즐 상태 점검\n관수량 조절 후 잔여 수분 확인'},
    {'step': '물주기', 'content': '이끼 관수 및 엽면 수분 보충\n배지 하부 수분 침투 확인\n관수 후 통풍 관리'},

    # ---- 하우스관리 관련 (C) ----
    {'step': '하우스관리', 'content': '하우스 환기창 개폐 조절\n내부 온도 및 습도 측정 기록\n이끼 건조 여부 점검 후 조치'},
    {'step': '하우스관리', 'content': '하우스 보온 자재 상태 점검\n시설 노후 부분 보수 작업\n이끼 재배 환경 온습도 관리'},
    {'step': '하우스관리', 'content': '하우스 내부 청소 및 정리\n환기팬 작동 상태 확인\n차광막 조절로 광량 관리'},
    {'step': '하우스관리', 'content': '하우스 측면 비닐 점검 및 보수\n난방기 가동 상태 확인\n이끼 재배상 온도 분포 확인'},

    # ---- 기타작업 관련 (D) ----
    {'step': '기타작업', 'content': '이끼 주변 잡초 제거 작업\n이끼 모듬 정리 및 분리\n작업 후 관수 실시'},
    {'step': '기타작업', 'content': '이끼 포장 상태 정돈\n불필요한 잔재물 수거 및 처리\n재배상 바닥 청소'},
    {'step': '기타작업', 'content': '재배 도구 세척 및 정비\n이끼 수확물 선별 및 포장 준비\n작업장 주변 정리정돈'},
    {'step': '기타작업', 'content': '이끼 재배 자재 정리 및 재고 확인\n배지 추가 준비 작업\n작업 일지 정리 및 기록'},

    # ---- 예찰활동 관련 (E) ----
    {'step': '예찰활동', 'content': '이끼 생육 상태 전체 점검\n약한 개체 분리 및 격리\n재배 밀도 조절 작업'},
    {'step': '예찰활동', 'content': '이끼 병해충 예찰 순회\n이상 개체 표식 및 기록\n필요 시 관리 계획 수립'},
    {'step': '예찰활동', 'content': '이끼 색상 및 수분 상태 관찰\n포트별 성장 속도 비교 기록\n환경 스트레스 징후 확인'},
    {'step': '예찰활동', 'content': '이끼 재배상 전체 순회 점검\n곰팡이 발생 여부 확인\n통풍 상태 점검 후 개선 조치'},
]

WMO_TO_KOREAN = {
    0: '맑음', 1: '구름조금', 2: '구름조금', 3: '구름많음',
    45: '흐림', 48: '흐림',
    51: '비', 53: '비', 55: '비', 56: '비', 57: '비',
    61: '비', 63: '비', 65: '비',
    66: '눈/비', 67: '눈/비',
    71: '눈', 73: '눈', 75: '눈', 77: '눈',
    80: '비', 81: '비', 82: '비',
    85: '눈', 86: '눈',
    95: '비', 96: '비', 99: '비',
}

DAY_NAMES = ['월', '화', '수', '목', '금', '토', '일']

# ============================================================
# 날씨
# ============================================================

def fetch_weather(start_date, end_date):
    """Open-Meteo API로 고양시 대화동 실제 날씨 데이터 조회"""
    params = {
        'latitude': LATITUDE,
        'longitude': LONGITUDE,
        'start_date': start_date,
        'end_date': end_date,
        'daily': 'temperature_2m_max,temperature_2m_min,precipitation_sum,relative_humidity_2m_mean,weathercode',
        'timezone': 'Asia/Seoul',
    }
    url = f'https://archive-api.open-meteo.com/v1/archive?{urlencode(params)}'
    print(f'  날씨 데이터 조회 중... ({start_date} ~ {end_date})')

    try:
        ctx = ssl.create_default_context()
        ctx.check_hostname = False
        ctx.verify_mode = ssl.CERT_NONE
        req = Request(url, headers={'User-Agent': 'farm-diary-generator/1.0'})
        with urlopen(req, timeout=30, context=ctx) as resp:
            data = json.loads(resp.read().decode())

        weather_map = {}
        daily = data.get('daily', {})
        dates = daily.get('time', [])
        max_temps = daily.get('temperature_2m_max', [])
        min_temps = daily.get('temperature_2m_min', [])
        precips = daily.get('precipitation_sum', [])
        humidities = daily.get('relative_humidity_2m_mean', [])
        codes = daily.get('weathercode', [])

        for i, date in enumerate(dates):
            wmo = codes[i] if i < len(codes) else 0
            weather_map[date] = {
                'max_temp': round(max_temps[i], 1) if i < len(max_temps) and max_temps[i] is not None else None,
                'min_temp': round(min_temps[i], 1) if i < len(min_temps) and min_temps[i] is not None else None,
                'precipitation': round(precips[i], 1) if i < len(precips) and precips[i] is not None else None,
                'humidity': round(humidities[i]) if i < len(humidities) and humidities[i] is not None else None,
                'weather': WMO_TO_KOREAN.get(wmo, '흐림'),
            }
        print(f'  ✓ {len(weather_map)}일간 실제 날씨 데이터 수신')
        return weather_map
    except Exception as e:
        print(f'  ⚠ 날씨 API 실패: {e}')
        print(f'  → 월별 기준 데이터로 대체')
        return None


def get_fallback_weather(date_str):
    month = int(date_str.split('-')[1])
    ranges = {
        1: {'low': (-8, -2), 'high': (1, 6), 'hum': (40, 55)},
        2: {'low': (-5, 1), 'high': (3, 9), 'hum': (40, 55)},
        3: {'low': (-1, 5), 'high': (8, 16), 'hum': (40, 55)},
        4: {'low': (5, 10), 'high': (15, 22), 'hum': (45, 60)},
        5: {'low': (10, 16), 'high': (20, 27), 'hum': (50, 65)},
        6: {'low': (17, 22), 'high': (25, 30), 'hum': (60, 75)},
        7: {'low': (22, 25), 'high': (28, 33), 'hum': (70, 85)},
        8: {'low': (22, 25), 'high': (28, 33), 'hum': (70, 85)},
        9: {'low': (15, 20), 'high': (23, 28), 'hum': (55, 70)},
        10: {'low': (7, 12), 'high': (16, 22), 'hum': (45, 60)},
        11: {'low': (0, 5), 'high': (8, 14), 'hum': (45, 55)},
        12: {'low': (-6, 0), 'high': (1, 7), 'hum': (40, 55)},
    }
    r = ranges.get(month, ranges[1])
    weathers = ['맑음', '구름조금', '구름많음', '흐림']
    return {
        'max_temp': round(random.uniform(*r['high']), 1),
        'min_temp': round(random.uniform(*r['low']), 1),
        'precipitation': None,
        'humidity': random.randint(*r['hum']),
        'weather': random.choice(weathers),
    }

# ============================================================
# 스케줄 생성
# ============================================================

def generate_schedule(start_date, end_date, freq=6, holidays=None):
    """작업 스케줄 생성"""
    if holidays is None:
        holidays = HOLIDAYS_2026

    start = datetime.strptime(start_date, '%Y-%m-%d')
    end = datetime.strptime(end_date, '%Y-%m-%d')

    # 작업 가능일 수집
    workdays = []
    current = start
    while current <= end:
        date_str = current.strftime('%Y-%m-%d')
        weekday = current.weekday()  # 0=월 ~ 6=일

        if freq >= 6:
            # 주 6일: 일요일(6)만 제외
            is_workday = weekday < 6
        elif freq == 5:
            # 주 5일: 토(5), 일(6) 제외
            is_workday = weekday < 5
        else:
            is_workday = weekday < 5

        if is_workday and date_str not in holidays:
            workdays.append({'date': date_str, 'day': DAY_NAMES[weekday], 'weekday': weekday})

        current += timedelta(days=1)

    # 주 3회 이하면 선별
    if freq <= 3:
        workdays = select_mwf_pattern(workdays)
    elif freq == 4:
        workdays = select_every_n_pattern(workdays, 4)

    # 패턴 적용 (셔플로 자연스럽게)
    schedule = []
    pattern_pool = list(range(len(PATTERNS)))
    random.seed(42)  # 재현 가능한 랜덤
    random.shuffle(pattern_pool)

    for i, day in enumerate(workdays):
        p_idx = pattern_pool[i % len(pattern_pool)]
        p = PATTERNS[p_idx]

        schedule.append({
            'no': i + 1,
            'date': day['date'],
            'day': day['day'],
            'pattern': p['step'][:1],
            'step': p['step'],
            'content': p['content'],
        })

    return schedule


def select_mwf_pattern(weekdays):
    selected = []
    last_date = None
    for day in weekdays:
        if last_date is None:
            selected.append(day)
            last_date = datetime.strptime(day['date'], '%Y-%m-%d')
        else:
            current = datetime.strptime(day['date'], '%Y-%m-%d')
            if (current - last_date).days >= 2:
                selected.append(day)
                last_date = current
    return selected


def select_every_n_pattern(weekdays, target_per_week):
    selected = []
    week_start = None
    week_count = 0
    for day in weekdays:
        current = datetime.strptime(day['date'], '%Y-%m-%d')
        week_num = current.isocalendar()[1]
        if week_start != week_num:
            week_start = week_num
            week_count = 0
        if week_count < target_per_week:
            selected.append(day)
            week_count += 1
    return selected

# ============================================================
# 엑셀
# ============================================================

def find_template():
    patterns = [
        os.path.expanduser('~/Downloads/영농일지+일괄등록*.xlsx'),
        os.path.expanduser('~/Downloads/영농일지*일괄등록*.xlsx'),
    ]
    for pattern in patterns:
        files = glob.glob(pattern)
        if files:
            originals = [f for f in files if '완성본' not in f]
            if originals:
                originals.sort(key=os.path.getmtime, reverse=True)
                return originals[0]
            files.sort(key=os.path.getmtime, reverse=True)
            return files[0]
    return None


def find_pilji_in_template(ws2, target_keyword='대화동'):
    for row_num in range(1, ws2.max_row + 1):
        val = ws2.cell(row=row_num, column=1).value
        if val and target_keyword in str(val):
            return val, row_num
    return None, None


def find_pumjong_options(ws2, pilji_row):
    options = []
    for col in range(2, ws2.max_column + 1):
        val = ws2.cell(row=pilji_row, column=col).value
        if val:
            options.append(val)
    return options


def validate_task_steps(ws3, pumjong_name, needed_steps):
    task_row = None
    for row_num in range(1, ws3.max_row + 1):
        val = ws3.cell(row=row_num, column=1).value
        if val == pumjong_name:
            task_row = row_num
            break
    if not task_row:
        return False, needed_steps

    available = []
    for col in range(2, ws3.max_column + 1):
        val = ws3.cell(row=task_row, column=col).value
        if val:
            available.append(val)

    missing = [s for s in needed_steps if s not in available]
    return len(missing) == 0, missing


def fill_excel(template_path, output_path, schedule, weather_data, pilji_name, pumjong_name, public='아니오'):
    import openpyxl

    wb = openpyxl.load_workbook(template_path)
    ws = wb['작업일지']

    # 기존 데이터 클리어
    for row in range(3, ws.max_row + 1):
        for col in range(1, 12):
            ws.cell(row=row, column=col).value = None

    needed_rows = len(schedule)
    max_data_row = 2 + needed_rows

    # 추가 행 서식 복사
    for extra_row in range(ws.max_row + 1, max_data_row + 1):
        for col in range(1, 12):
            src_cell = ws.cell(row=3, column=col)
            dst_cell = ws.cell(row=extra_row, column=col)
            if src_cell.has_style:
                dst_cell.font = copy.copy(src_cell.font)
                dst_cell.border = copy.copy(src_cell.border)
                dst_cell.fill = copy.copy(src_cell.fill)
                dst_cell.number_format = src_cell.number_format
                dst_cell.protection = copy.copy(src_cell.protection)
                dst_cell.alignment = copy.copy(src_cell.alignment)

    # 데이터 입력
    for i, entry in enumerate(schedule):
        row = i + 3
        parts = entry['date'].split('-')
        dt = datetime(int(parts[0]), int(parts[1]), int(parts[2]))

        wd = (weather_data or {}).get(entry['date']) or get_fallback_weather(entry['date'])

        ws.cell(row=row, column=1).value = dt
        ws.cell(row=row, column=1).number_format = 'mm-dd-yy'

        for col_idx, val in [(2, pilji_name), (3, pumjong_name), (4, entry['step']),
                              (5, entry['content']), (6, wd['weather'])]:
            cell = ws.cell(row=row, column=col_idx)
            cell.value = val
            cell.number_format = '@'

        if wd.get('min_temp') is not None:
            ws.cell(row=row, column=7).value = wd['min_temp']
        if wd.get('max_temp') is not None:
            ws.cell(row=row, column=8).value = wd['max_temp']
        if wd.get('precipitation') and wd['precipitation'] > 0:
            ws.cell(row=row, column=9).value = wd['precipitation']
        if wd.get('humidity') is not None:
            ws.cell(row=row, column=10).value = wd['humidity']

        cell_k = ws.cell(row=row, column=11)
        cell_k.value = public
        cell_k.number_format = '@'

    # 데이터 유효성 범위 확장
    for dv in ws.data_validations.dataValidation:
        sqref = str(dv.sqref)
        for col_letter in ['A', 'D', 'F', 'K']:
            old = f'{col_letter}3:{col_letter}33'
            new = f'{col_letter}3:{col_letter}{max_data_row}'
            if old in sqref:
                dv.sqref = sqref.replace(old, new)
                sqref = str(dv.sqref)

    wb.save(output_path)
    return len(schedule)


# ============================================================
# 메인
# ============================================================

def main():
    parser = argparse.ArgumentParser(description='agrion.kr 영농일지 일괄등록 엑셀 생성기')
    parser.add_argument('--start', required=True, help='시작일 (YYYY-MM-DD)')
    parser.add_argument('--end', required=True, help='종료일 (YYYY-MM-DD)')
    parser.add_argument('--freq', type=int, default=6, help='주당 작업일 수 (기본: 6, 일요일만 제외)')
    parser.add_argument('--template', help='엑셀 템플릿 경로')
    parser.add_argument('--output', help='출력 파일 경로')
    parser.add_argument('--pilji', default='대화동', help='필지 키워드 (기본: 대화동)')
    parser.add_argument('--pumjong', default='기타', help='품종명 (기본: 기타)')
    parser.add_argument('--public', default='아니오', help='공개여부 (기본: 아니오)')
    args = parser.parse_args()

    print('=' * 55)
    print('  agrion.kr 영농일지 일괄등록 엑셀 생성기')
    print('=' * 55)
    print(f'  기간: {args.start} ~ {args.end}')
    print(f'  빈도: 주 {args.freq}일')
    print()

    # 1. 템플릿
    template_path = args.template or find_template()
    if not template_path or not os.path.exists(template_path):
        print('❌ 엑셀 템플릿을 찾을 수 없습니다.')
        print('   agrion.kr에서 "영농일지 일괄등록" 엑셀을 다운로드 후 --template으로 지정하세요.')
        sys.exit(1)
    print(f'[1] 템플릿: {os.path.basename(template_path)}')

    # 2. 필지/품종 확인
    import openpyxl
    wb_check = openpyxl.load_workbook(template_path, data_only=True)
    ws2 = wb_check['Data2']
    ws3 = wb_check['Data3']

    pilji_name, pilji_row = find_pilji_in_template(ws2, args.pilji)
    if not pilji_name:
        print(f'❌ 필지 "{args.pilji}"를 찾을 수 없습니다.')
        print('   등록된 필지:')
        for r in range(1, ws2.max_row + 1):
            v = ws2.cell(row=r, column=1).value
            if v:
                print(f'     - {v}')
        sys.exit(1)
    print(f'[2] 필지: {pilji_name}')

    pumjong_options = find_pumjong_options(ws2, pilji_row)
    pumjong_name = args.pumjong if args.pumjong in pumjong_options else (pumjong_options[0] if pumjong_options else args.pumjong)
    print(f'[3] 품종: {pumjong_name} (옵션: {pumjong_options})')

    needed = list(set(p['step'] for p in PATTERNS))
    valid, missing = validate_task_steps(ws3, pumjong_name, needed)
    if valid:
        print(f'[4] 작업단계: 모두 유효 ✓')
    else:
        print(f'⚠ 작업단계 누락: {missing}')

    wb_check.close()

    # 3. 스케줄
    schedule = generate_schedule(args.start, args.end, args.freq)
    print(f'[5] 스케줄: {len(schedule)}건 생성')

    # 4. 날씨
    weather_data = fetch_weather(args.start, args.end)

    # 5. 엑셀 생성
    if args.output:
        output_path = args.output
    else:
        end_month = args.end[:7].replace('-', '')
        output_path = os.path.expanduser(f'~/Downloads/영농일지_일괄등록_{end_month}.xlsx')

    count = fill_excel(template_path, output_path, schedule, weather_data, pilji_name, pumjong_name, args.public)

    print()
    print('=' * 55)
    print(f'  ✓ 완료! {count}건 작업일지 생성')
    print(f'  파일: {output_path}')
    print('=' * 55)

    # 미리보기
    print('\n=== 미리보기 (처음 10건 / 마지막 5건) ===')
    preview = schedule[:10] + (['...'] if len(schedule) > 15 else []) + schedule[-5:]
    for entry in preview:
        if entry == '...':
            print(f'  ... ({len(schedule) - 15}건 생략) ...')
            continue
        wd = (weather_data or {}).get(entry['date']) or get_fallback_weather(entry['date'])
        temp = f"{wd['min_temp']}~{wd['max_temp']}℃" if wd.get('min_temp') is not None else '?'
        preview_text = entry['content'].replace('\n', ' / ')[:50]
        print(f"  {entry['no']:3d}. {entry['date']} ({entry['day']}) | {entry['step']:6s} | {wd['weather']:4s} {temp:15s} | {preview_text}")


if __name__ == '__main__':
    main()
