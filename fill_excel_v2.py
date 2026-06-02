import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from datetime import datetime, timedelta
import json
import copy
import random

# 원본 템플릿 읽기
src = '/Downloads/영농일지+일괄등록_20260323.xlsx'
dst = '/Downloads/영농일지_일괄등록_완성본_v2.xlsx'

wb = openpyxl.load_workbook(src)
ws = wb['작업일지']

# 스케줄 데이터 읽기
with open('/.omc/scientist/work_schedule.json', 'r') as f:
    schedule = json.load(f)

# Data2에서 정확한 필지명 확인
ws2 = wb['Data2']
pilji_name = None
for row in ws2.iter_rows(min_col=1, max_col=1, values_only=False):
    for cell in row:
        if cell.value and '대화동' in str(cell.value):
            pilji_name = cell.value
            pilji_row = cell.row
            break

print(f'필지명 확인: "{pilji_name}" (Data2 Row {pilji_row})')

# 해당 필지의 품종 목록 확인
pumjong_options = []
for col in range(2, ws2.max_column + 1):
    val = ws2.cell(row=pilji_row, column=col).value
    if val:
        pumjong_options.append(val)
print(f'품종 옵션: {pumjong_options}')

# 품종 = "기타" 사용
pumjong_name = '기타'
if '기타' in pumjong_options:
    pumjong_name = '기타'
elif pumjong_options:
    pumjong_name = pumjong_options[0]
print(f'선택 품종: "{pumjong_name}"')

# Data3에서 해당 품종의 작업단계 목록 확인
ws3 = wb['Data3']
task_options = []
task_row = None
for row in ws3.iter_rows(min_col=1, max_col=1, values_only=False):
    for cell in row:
        if cell.value == pumjong_name:
            task_row = cell.row
            break

if task_row:
    for col in range(2, ws3.max_column + 1):
        val = ws3.cell(row=task_row, column=col).value
        if val:
            task_options.append(val)

print(f'작업단계 옵션 ({len(task_options)}개): {task_options[:10]}...')

# 우리가 사용할 작업단계가 유효한지 확인
needed_tasks = ['파종', '물주기', '하우스관리', '기타작업', '예찰활동']
for task in needed_tasks:
    if task in task_options:
        print(f'  ✓ "{task}" - 유효')
    else:
        print(f'  ✗ "{task}" - 없음! 대체값 필요')

# 날씨 옵션 확인
ws_data = wb['Data']
weather_options = []
for row in ws_data.iter_rows(min_col=1, max_col=1, values_only=True):
    if row[0]:
        weather_options.append(row[0])
print(f'날씨 옵션: {weather_options}')

# 월별 날씨/기온 생성
def get_weather(date_str):
    month = int(date_str.split('-')[1])
    ranges = {
        1: {'low': (-8, -2), 'high': (1, 6), 'hum': (40, 55)},
        2: {'low': (-5, 1), 'high': (3, 9), 'hum': (40, 55)},
        3: {'low': (-1, 5), 'high': (8, 16), 'hum': (40, 55)},
    }
    r = ranges[month]

    # 겨울~초봄: 맑음/구름조금/구름많음/흐림 위주
    winter_weather = ['맑음', '구름조금', '구름많음', '흐림']
    w = random.choice(winter_weather)

    low = round(random.uniform(*r['low']), 1)
    high = round(random.uniform(*r['high']), 1)
    hum = random.randint(*r['hum'])

    return w, low, high, hum

# Row 3 셀 서식 복사를 위한 함수
def copy_cell_format(src_cell, dst_cell):
    """셀 서식만 복사"""
    if src_cell.has_style:
        dst_cell.font = copy.copy(src_cell.font)
        dst_cell.border = copy.copy(src_cell.border)
        dst_cell.fill = copy.copy(src_cell.fill)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy.copy(src_cell.protection)
        dst_cell.alignment = copy.copy(src_cell.alignment)

# 기존 데이터 클리어 (Row 3 ~ Row 33)
print('\n기존 데이터 클리어 중...')
for row in range(3, 34):
    for col in range(1, 12):
        ws.cell(row=row, column=col).value = None

# 32건 데이터가 필요하므로 Row 34도 필요
# Row 34 추가 (Row 3 서식 복사)
if len(schedule) > 31:
    print(f'32건 데이터: Row 34 추가 필요')
    for col in range(1, 12):
        src_cell = ws.cell(row=3, column=col)
        dst_cell = ws.cell(row=34, column=col)
        copy_cell_format(src_cell, dst_cell)

# 데이터 입력
print('\n데이터 입력 중...')
for i, entry in enumerate(schedule):
    row = i + 3  # Row 3부터 시작

    # 날짜 (datetime 객체로)
    date_parts = entry['date'].split('-')
    dt = datetime(int(date_parts[0]), int(date_parts[1]), int(date_parts[2]))

    # 날씨 데이터
    weather, low_temp, high_temp, humidity = get_weather(entry['date'])

    # 작업내용에서 <br>을 줄바꿈으로
    content = entry['content'].replace('<br>', '\n')

    # 셀에 값 입력
    cell_a = ws.cell(row=row, column=1)  # 일자
    cell_a.value = dt
    cell_a.number_format = 'mm-dd-yy'

    cell_b = ws.cell(row=row, column=2)  # 필지
    cell_b.value = pilji_name
    cell_b.number_format = '@'

    cell_c = ws.cell(row=row, column=3)  # 품종
    cell_c.value = pumjong_name
    cell_c.number_format = '@'

    cell_d = ws.cell(row=row, column=4)  # 작업단계
    cell_d.value = entry['step']
    cell_d.number_format = '@'

    cell_e = ws.cell(row=row, column=5)  # 작업내용
    cell_e.value = content
    cell_e.number_format = '@'

    cell_f = ws.cell(row=row, column=6)  # 날씨
    cell_f.value = weather
    cell_f.number_format = '@'

    cell_g = ws.cell(row=row, column=7)  # 최저온도
    cell_g.value = low_temp

    cell_h = ws.cell(row=row, column=8)  # 최고온도
    cell_h.value = high_temp

    # 강수량 (비인 경우만)
    cell_i = ws.cell(row=row, column=9)
    if weather == '비':
        cell_i.value = round(random.uniform(1, 10), 1)

    cell_j = ws.cell(row=row, column=10)  # 습도
    cell_j.value = humidity
    cell_j.number_format = '0_);[Red]\\(0\\)'

    cell_k = ws.cell(row=row, column=11)  # 공개여부
    cell_k.value = '아니오'
    cell_k.number_format = '@'

    print(f'  {i+1}. {entry["date"]} ({entry["day"]}) | {pilji_name[-10:]} | {pumjong_name} | {entry["step"]} | {weather} {low_temp}~{high_temp}℃')

# 데이터 유효성 검사 범위 확장 (Row 34 포함)
# 기존 validation은 유지하되, D/F/K가 Row 33까지만이면 확장
print('\n데이터 유효성 범위 확인 및 조정...')
for dv in ws.data_validations.dataValidation:
    sqref = str(dv.sqref)
    if 'D3:D33' in sqref:
        dv.sqref = 'D3:D34'
        print(f'  작업단계 범위 확장: D3:D33 → D3:D34')
    if 'F3:F33' in sqref:
        dv.sqref = 'F3:F34'
        print(f'  날씨 범위 확장: F3:F33 → F3:F34')
    if 'K3:K33' in sqref:
        dv.sqref = 'K3:K34'
        print(f'  공개여부 범위 확장: K3:K33 → K3:K34')
    if 'A3:A33' in sqref:
        dv.sqref = 'A3:A34'
        print(f'  일자 범위 확장: A3:A33 → A3:A34')

# 저장
wb.save(dst)
print(f'\n{"="*50}')
print(f'파일 저장 완료: {dst}')
print(f'총 {len(schedule)}건 작업일지 입력 완료')
print(f'{"="*50}')

# 검증
print('\n=== 검증 ===')
wb2 = openpyxl.load_workbook(dst)
ws2_check = wb2['작업일지']
count = 0
for row in range(3, 35):
    val = ws2_check.cell(row=row, column=1).value
    if val:
        count += 1
        date_str = val.strftime('%Y-%m-%d') if isinstance(val, datetime) else str(val)
        pilji = ws2_check.cell(row=row, column=2).value
        pumjong = ws2_check.cell(row=row, column=3).value
        task = ws2_check.cell(row=row, column=4).value
        print(f'  {count}. {date_str} | {pilji} | {pumjong} | {task}')
print(f'\n총 {count}건 확인')
